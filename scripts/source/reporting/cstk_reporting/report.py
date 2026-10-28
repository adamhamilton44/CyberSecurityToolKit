"""Deterministic editable report generation. Uses no AI and makes no network calls."""

from __future__ import annotations

import hashlib
import json
import os
import signal
import subprocess
import tempfile
from collections import Counter
from pathlib import Path

from .model import SEVERITIES, readiness
from .storage import (
    NotebookError,
    canonical,
    digest,
    now,
    private_dir,
)

REPORT_FILES = ("report.docx", "report.pdf", "report.md", "manifest.json")


def report_blocks(state, status, report_id, reviewer="", finalized_at=""):
    profile = state["profile"]
    scope = state["scope_history"][-1]
    answer = lambda key: profile.get(key) or "Not recorded"
    confirmed = sorted(
        (f for f in state["findings"] if f["status"] == "confirmed"),
        key=lambda f: (SEVERITIES.index(f["severity"]), f["id"]),
    )
    counts = Counter(f["severity"] for f in confirmed)
    complete = sum(run["outcome"] == "completed" for run in state["activities"])
    blocks = []

    def add(kind, text=None, **kwargs):
        blocks.append(dict(kind=kind, text=text, **kwargs))

    add("title", "Security Assessment Report")
    add("subtitle", answer("client"))
    add(
        "p",
        f"This report documents {answer('engagement_name')} for {answer('client')}. It presents confirmed findings, supporting evidence, and coverage limitations for remediation planning.",
    )
    add(
        "table",
        headers=["Report detail", "Value"],
        rows=[
            ["Status", status + (" | " + finalized_at if finalized_at else "")],
            ["Engagement reference", answer("engagement_reference")],
            ["Report reference", report_id],
            [
                "Assessment period",
                answer("testing_start") + " to " + answer("testing_end"),
            ],
            [
                "Prepared and reviewed by",
                answer("author") + (" / " + reviewer if reviewer else ""),
            ],
        ],
        widths=[1.65, 5.25],
    )
    add("h1", "Executive summary")
    if profile.get("executive_context"):
        add("p", profile["executive_context"])
    add(
        "p",
        f"The assessment record contains {len(confirmed)} confirmed finding{'s' if len(confirmed) != 1 else ''} and {len(state['activities'])} recorded assessment activit{'ies' if len(state['activities']) != 1 else 'y'}, of which {complete} completed. Findings describe the evidence available during the stated assessment period; they do not establish the security of every asset or application route.",
    )
    if not confirmed:
        add(
            "p",
            "No findings have been confirmed in the captured evidence. This does not establish that vulnerabilities are absent; review the coverage and limitations before relying on this result.",
        )
    if confirmed:
        add(
            "p",
            "Confirmed findings by severity: "
            + " | ".join(f"{severity} {counts[severity]}" for severity in SEVERITIES)
            + ".",
        )
    if status != "Final":
        pending = sum(f["status"] == "candidate" for f in state["findings"])
        errors = readiness(state)
        add(
            "p",
            f"Draft review status: {pending} unresolved finding candidate(s) and {len(errors)} reporting readiness issue(s). This document has not been finalized.",
        )
    add("h1", "Findings summary")
    if confirmed:
        add(
            "table",
            headers=["Finding reference", "Severity", "Finding", "Affected asset"],
            rows=[[f["id"], f["severity"], f["title"], f["asset"]] for f in confirmed],
            widths=[1.25, 0.8, 2.85, 2.0],
        )
    else:
        add("p", "There are no confirmed findings to summarize.")
    dispositions = Counter(f["status"] for f in state["findings"])
    add(
        "p",
        f"Review dispositions: {dispositions['false_positive']} false positive(s), {dispositions['duplicate']} duplicate(s), {dispositions['excluded']} excluded observation(s). Their internal evidence and notes are not included as confirmed findings.",
    )
    if status != "Final" and dispositions["candidate"]:
        add("h2", "Candidates requiring review")
        add(
            "table",
            headers=["Reference", "Candidate", "Proposed severity"],
            rows=[
                [f["id"], f["title"], f["severity"]]
                for f in state["findings"]
                if f["status"] == "candidate"
            ],
            widths=[1.4, 4.4, 1.1],
        )
    add("page")
    add("h1", "Scope and assessment context")
    add("p", answer("business_context"))
    add("h2", "Objectives and priority assets")
    add("p", answer("objectives"))
    if profile.get("priority_assets"):
        add("p", profile["priority_assets"])
    add("h2", "Recorded scope")
    add(
        "p",
        f"Scope revision {scope['revision']} is a reporting snapshot of engagement {state['engagement_id']}. The snapshot records the approved boundaries supplied by the operator; it does not authorize or expand testing.",
    )
    add(
        "p",
        "Authorization reference: "
        + (scope.get("authorization_reference") or "Not recorded"),
    )
    add(
        "table",
        headers=["Boundary", "Type", "Asset"],
        rows=[
            [label, entry["type"], entry["value"]]
            for key, label in (("include", "Included"), ("exclude", "Excluded"))
            for entry in scope["scope"][key]
        ]
        or [["Not recorded", "", ""]],
        widths=[1.0, 1.0, 4.9],
    )
    add("h2", "Methodology")
    add("p", answer("assessment_type"))
    add("p", answer("methodology"))
    add("h2", "Coverage limitations")
    add("p", answer("limitations"))
    add(
        "p",
        "Coverage is limited to the activities and evidence recorded in this notebook. A successful connection is not a vulnerability assessment. Errors, partial results, and skipped checks do not constitute passing security tests. The report does not certify compliance with a standard.",
    )
    for finding in confirmed:
        add("page")
        add("h1", finding["title"])
        add(
            "p",
            f"{finding['id']}   |   {finding['severity']}   |   {finding['asset']}",
            bold=True,
        )
        add(
            "p",
            "Recorded scope revision: "
            + str(finding.get("scope_revision", "Not recorded")),
        )
        for heading, key in (
            ("Observation", "observation"),
            ("Business and technical impact", "impact"),
            ("Severity rationale", "severity_reason"),
            ("Validation", "validation"),
            ("Remediation", "remediation"),
        ):
            add("h2", heading)
            add("p", finding[key])
        if finding.get("owner") or finding.get("target_date"):
            add(
                "p",
                f"Remediation owner: {finding.get('owner') or 'To be assigned'}. Target date: {finding.get('target_date') or 'To be agreed'}.",
            )
        if finding.get("references"):
            add("h2", "References")
            add("p", finding["references"])
        add("h2", "Supporting evidence")
        for evidence_id in finding["evidence_ids"]:
            note = next(n for n in state["notes"] if n["id"] == evidence_id)
            add("p", f"{note['id']}   {note['title']}   {note['at']}", bold=True)
            body = note["body"]
            excerpt = body[:1800]
            if len(body) > 1800:
                excerpt += "\n[Excerpt limited to 1800 characters. Full note remains in the encrypted notebook.]"
            add("evidence", excerpt)
        add(
            "p",
            f"Finding reviewed by {finding['reviewer']} at {finding['reviewed_at']}. Severity is an operator assessment; no CVSS score has been inferred.",
        )
    add("page")
    add("h1", "Remediation plan")
    if confirmed:
        add(
            "table",
            headers=["Finding", "Severity", "Owner", "Target date"],
            rows=[
                [
                    f["id"] + " " + f["title"],
                    f["severity"],
                    f.get("owner") or "To be assigned",
                    f.get("target_date") or "To be agreed",
                ]
                for f in confirmed
            ],
            widths=[3.4, 0.8, 1.5, 1.2],
        )
        add(
            "p",
            "Validate changes in an appropriate environment, preserve supporting evidence, and schedule retesting for the affected assets. Closure requires new validation; generating this report does not mark a finding as fixed.",
        )
    else:
        add(
            "p",
            "No confirmed findings have remediation actions recorded. Outstanding coverage gaps should be addressed before broader security conclusions are drawn.",
        )
    if profile.get("remediation_context"):
        add("p", profile["remediation_context"])
    notes = [
        n
        for n in state["notes"]
        if n["report_include"] and not n["withdrawn"] and n["kind"] != "evidence"
    ]
    if notes:
        add("h2", "Additional assessment context")
        for note in notes:
            add("h2", note["title"])
            add("p", note["body"])
    add("page")
    add("h1", "Assessment activity and handling")
    add(
        "table",
        headers=["Tool and origin", "Target", "Outcome", "Captured completion time"],
        rows=[
            [
                a["tool"]
                + "\n"
                + ("Adapter capture" if a["origin"] == "adapter" else "Operator entry"),
                a["target"],
                a["outcome"],
                a["finished_at"] or "Not completed",
            ]
            for a in state["activities"]
        ]
        or [["No activities recorded", "", "", ""]],
        widths=[1.5, 2.5, 0.85, 2.05],
    )
    for activity in state["activities"]:
        add(
            "p",
            f"{activity['id']} | {activity['tool']} | {activity['outcome']}: {activity['summary'] or 'No completion summary recorded.'}",
        )
    add("h2", "Reporting and evidence handling")
    add("p", "Intended audience: " + answer("audience"))
    add("p", "Approved recipients: " + answer("recipients"))
    add("p", "Delivery method: " + answer("delivery_method"))
    add("p", answer("evidence_policy"))
    add(
        "p",
        f"Notebook retention: {answer('retention_days')} days after finalization, with explicit operator action required for deletion. Legal hold: {answer('legal_hold')}. Final report files are separate deliverables and are not encrypted by the notebook.",
    )
    if profile.get("escalation_reference"):
        add("p", "Critical finding coordination: " + profile["escalation_reference"])
    add("h2", "Report basis")
    add(
        "p",
        "Recorded scope revisions: "
        + ", ".join(str(s["revision"]) for s in state["scope_history"]),
    )
    add("p", "Current scope SHA256: " + scope["scope_sha256"])
    add(
        "p",
        "File hashes in the accompanying manifest identify this export. Edits to the Word document or PDF create a different artifact and require renewed review. Captured adapter and operator records are reporting inputs, not independent proof of authorization or execution.",
    )
    return blocks


def _markdown(blocks):
    lines = []
    for block in blocks:
        kind = block["kind"]
        value = block.get("text")
        if kind in ("title", "h1"):
            lines += ["# " + value, ""]
        elif kind in ("subtitle", "h2"):
            lines += ["## " + value, ""]
        elif kind in ("p", "evidence"):
            lines += [value, ""]
        elif kind == "table":

            def row(items):
                return (
                    "| "
                    + " | ".join(
                        str(x).replace("|", "/").replace("\n", " ") for x in items
                    )
                    + " |"
                )

            lines += [
                row(block["headers"]),
                row(["---"] * len(block["headers"])),
                *[row(items) for items in block["rows"]],
                "",
            ]
    return "\n".join(lines)


def _docx(blocks, path, classification, reference):
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    # The bundled default Word template carries a blue Title border; remove it.
    for style in document.styles:
        properties = style.element.find(qn("w:pPr"))
        if properties is not None:
            for border in list(properties.findall(qn("w:pBdr"))):
                properties.remove(border)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.8)
    section.header_distance = section.footer_distance = Inches(0.3)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in [
        ("Title", 27),
        ("Subtitle", 14),
        ("Heading 1", 17),
        ("Heading 2", 12),
    ]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(
            10 if name.startswith("Heading") else 0
        )
        style.paragraph_format.space_after = Pt(7)
    for block in blocks:
        kind = block["kind"]
        value = block.get("text")
        if kind == "page":
            document.add_page_break()
        elif kind in ("title", "subtitle", "h1", "h2"):
            document.add_paragraph(
                "".join(c if c.isalnum() or c.isspace() else " " for c in value),
                style={
                    "title": "Title",
                    "subtitle": "Subtitle",
                    "h1": "Heading 1",
                    "h2": "Heading 2",
                }[kind],
            )
        elif kind in ("p", "evidence"):
            paragraph = document.add_paragraph(value)
            if block.get("bold"):
                for run in paragraph.runs:
                    run.bold = True
            if kind == "evidence":
                for run in paragraph.runs:
                    run.font.name = "DejaVu Sans Mono"
                    run.font.size = Pt(8.5)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(12)
        elif kind == "table":
            table = document.add_table(rows=1, cols=len(block["headers"]))
            table.autofit = False
            for col, width in zip(table.columns, block["widths"]):
                col.width = Inches(width)
            for cell, label, width in zip(
                table.rows[0].cells, block["headers"], block["widths"]
            ):
                cell.width = Inches(width)
                cell.text = label
            for row in block["rows"]:
                cells = table.add_row().cells
                for cell, label, width in zip(cells, row, block["widths"]):
                    cell.width = Inches(width)
                    cell.text = str(label)
            for i, row in enumerate(table.rows):
                props = row._tr.get_or_add_trPr()
                if i == 0:
                    repeat = OxmlElement("w:tblHeader")
                    props.append(repeat)
                # Rows may split when content is genuinely longer than a page.
                for cell in row.cells:
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cp = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    for edge in ("top", "left", "bottom", "right"):
                        element = OxmlElement("w:" + edge)
                        element.set(qn("w:val"), "single")
                        element.set(qn("w:sz"), "4")
                        element.set(qn("w:color"), "D9D9D9")
                        borders.append(element)
                    cp.append(borders)
                    margins = OxmlElement("w:tcMar")
                    for side in ("top", "left", "bottom", "right"):
                        element = OxmlElement("w:" + side)
                        element.set(qn("w:w"), "95")
                        element.set(qn("w:type"), "dxa")
                        margins.append(element)
                    cp.append(margins)
                    shade = OxmlElement("w:shd")
                    shade.set(
                        qn("w:fill"),
                        "243746" if i == 0 else ("F1F4F6" if i % 2 == 0 else "FFFFFF"),
                    )
                    cp.append(shade)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(2)
                        paragraph.paragraph_format.space_before = Pt(2)
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
                            run.font.bold = i == 0
                            run.font.color.rgb = RGBColor.from_string(
                                "FFFFFF" if i == 0 else "000000"
                            )
            document.add_paragraph().paragraph_format.space_after = Pt(0)
    header = section.header.paragraphs[0]
    header.text = classification + "  |  " + reference
    header.style = document.styles["Normal"]
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
    document.core_properties.title = "Security Assessment Report"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.subject = ""
    document.core_properties.comments = ""
    document.save(path)


def _pdf(docx_path, destination, soffice):
    if not soffice:
        raise NotebookError(
            "PDF generation requires an explicitly selected LibreOffice executable (--soffice)."
        )
    binary = Path(soffice)
    if (
        not binary.is_absolute()
        or not binary.is_file()
        or not os.access(binary, os.X_OK)
    ):
        raise NotebookError("Select an absolute executable LibreOffice path.")
    with tempfile.TemporaryDirectory(prefix=".render-", dir=destination) as temp:
        env = {
            k: v
            for k, v in os.environ.items()
            if k in ("PATH", "LANG", "LC_ALL", "FONTCONFIG_PATH", "FONTCONFIG_FILE")
        }
        env.setdefault("PATH", "/usr/bin:/bin")
        env.update(HOME=temp, TMPDIR=temp, TEMP=temp, TMP=temp)
        command = [
            str(binary),
            "-env:UserInstallation=" + Path(temp, "profile").as_uri(),
            "--headless",
            "--convert-to",
            "pdf:writer_pdf_Export",
            "--outdir",
            str(destination),
            str(docx_path),
        ]
        try:
            process = subprocess.Popen(
                command,
                env=env,
                stdin=subprocess.DEVNULL,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                start_new_session=True,
            )
            try:
                process.communicate(timeout=90)
            except subprocess.TimeoutExpired as error:
                try:
                    os.killpg(process.pid, signal.SIGKILL)
                except ProcessLookupError:
                    pass
                process.communicate()
                raise NotebookError(
                    "PDF renderer timed out; its process group was stopped and no report was finalized."
                ) from error
        except OSError as error:
            raise NotebookError(
                "PDF renderer failed; no report was finalized."
            ) from error
        pdf = destination / "report.pdf"
        if process.returncode or not pdf.is_file() or pdf.stat().st_size < 100:
            raise NotebookError(
                "PDF renderer did not create a valid report; no report was finalized."
            )


def export_bundle(
    state, directory, *, status, report_id, reviewer="", finalized_at="", soffice=None
):
    """New, private directory only. Failure removes its own partial plaintext files."""
    path = Path(os.path.abspath(os.path.expanduser(directory)))
    parent = path.parent
    if not parent.is_dir():
        raise NotebookError("Create the report output parent directory first.")
    # Anchor the output path and reject symlinked ancestors before creating it.
    # A private parent prevents another user replacing the newly created child.
    parent_fd = private_dir(parent)
    try:
        try:
            os.mkdir(path.name, 0o700, dir_fd=parent_fd)
        except FileExistsError as error:
            raise NotebookError(
                "Report output directory already exists; choose a new one."
            ) from error
    finally:
        os.close(parent_fd)
    fd = private_dir(path)
    old_umask = os.umask(0o077)
    try:
        blocks = report_blocks(state, status, report_id, reviewer, finalized_at)
        _docx(
            blocks,
            path / "report.docx",
            state["profile"].get("classification", "Classification not set"),
            report_id,
        )
        (path / "report.md").write_text(_markdown(blocks))
        _pdf(path / "report.docx", path, soffice)
        manifest = {
            "format": "cstk-report-1",
            "id": report_id,
            "status": status,
            "generator_version": "0.1.0",
            "created_at": now(),
            "report_model_sha256": digest(blocks),
            "files": {},
        }
        for name in REPORT_FILES[:3]:
            os.chmod(path / name, 0o600)
            manifest["files"][name] = hashlib.sha256(
                (path / name).read_bytes()
            ).hexdigest()
        (path / "manifest.json").write_bytes(canonical(manifest))
        os.chmod(path / "manifest.json", 0o600)
        for name in REPORT_FILES:
            check_fd = os.open(name, os.O_RDONLY | os.O_NOFOLLOW, dir_fd=fd)
            try:
                # Reports can exceed the notebook size limit; validate identity/mode separately.
                info = os.fstat(check_fd)
                if (
                    info.st_uid != os.geteuid()
                    or info.st_nlink != 1
                    or info.st_mode & 0o077
                ):
                    raise NotebookError("Unsafe report output file.")
                os.fsync(check_fd)
            finally:
                os.close(check_fd)
        os.fsync(fd)
        return manifest
    except Exception:
        for name in REPORT_FILES:
            try:
                os.unlink(name, dir_fd=fd)
            except FileNotFoundError:
                pass
        raise
    finally:
        os.umask(old_umask)
        os.close(fd)


def verify_bundle(directory, expected):
    fd = private_dir(directory)
    try:

        def read(name):
            f = os.open(name, os.O_RDONLY | os.O_NOFOLLOW, dir_fd=fd)
            try:
                import stat

                info = os.fstat(f)
                if (
                    not stat.S_ISREG(info.st_mode)
                    or info.st_uid != os.geteuid()
                    or info.st_nlink != 1
                    or info.st_mode & 0o077
                    or info.st_size > 64 * 1024 * 1024
                ):
                    raise NotebookError("Unsafe report file.")
                with os.fdopen(f, "rb", closefd=False) as stream:
                    return stream.read()
            finally:
                os.close(f)

        actual = json.loads(read("manifest.json"))
        if actual != expected:
            raise NotebookError("Report manifest changed after review/export.")
        if set(expected.get("files", {})) != set(REPORT_FILES[:3]):
            raise NotebookError("Report manifest has an unexpected file set.")
        for name, sha in expected["files"].items():
            if hashlib.sha256(read(name)).hexdigest() != sha:
                raise NotebookError("Report artifact changed after review/export.")
        return True
    except (FileNotFoundError, json.JSONDecodeError) as error:
        raise NotebookError("Report bundle is missing or incomplete.") from error
    finally:
        os.close(fd)


def discard_uncommitted_bundle(directory, expected):
    """Remove only the exact new export if notebook registration failed."""
    verify_bundle(directory, expected)
    fd = private_dir(directory)
    try:
        for name in REPORT_FILES:
            os.unlink(name, dir_fd=fd)
        os.fsync(fd)
    finally:
        os.close(fd)
    try:
        os.rmdir(directory)
    except OSError:
        pass
